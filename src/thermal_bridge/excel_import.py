"""Read the original workbook, including the formatting used by Psi."""

import argparse
import csv
from colorsys import rgb_to_hls, hls_to_rgb
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX


def _hex_color(color, theme, palette, default):
    if color is None or color.type == 'auto':
        return default
    if color.type == 'rgb':
        rgb = color.rgb[-6:]
    elif color.type == 'indexed' and color.indexed < len(palette):
        rgb = palette[color.indexed][-6:]
    elif color.type == 'theme' and color.theme < len(theme):
        rgb = theme[color.theme]
    else:
        raise ValueError(f'Unsupported Excel color: {color}')
    tint = color.tint
    channels = [int(rgb[i:i + 2], 16) for i in (0, 2, 4)]
    if tint:
        hue, lightness, saturation = rgb_to_hls(*(c / 255 for c in channels))
        lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
        channels = [round(c * 255) for c in hls_to_rgb(hue, lightness, saturation)]
    return '#' + ''.join(f'{c:02x}' for c in channels)


def read_excel_rows(filepath: str | Path, sheet_name: str = '44er Wand') -> list[dict]:
    """Return CSV-compatible rows. Colors come from the Bezeichnung cell.

    Formula values use Excel's saved cache; openpyxl does not calculate formulas
    or evaluate conditional formatting. Header rows 2 and 3 match this workbook.
    """
    workbook = load_workbook(filepath, data_only=True, read_only=True)
    try:
        sheet = workbook[sheet_name]
        theme = []
        if workbook.loaded_theme:
            root = ElementTree.fromstring(workbook.loaded_theme)
            namespace = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            scheme = root.find('.//a:clrScheme', namespace)
            for entry in scheme:
                color = entry[0]
                theme.append(color.get('lastClr') or color.get('val'))
        palette = workbook._colors or COLOR_INDEX
        # D2 is the group heading "berechnet durch"; D3:F3 are its columns.
        headers = [sheet.cell(3, col).value or sheet.cell(2, col).value
                   for col in range(1, 11)]
        if headers[0:2] != ['Bezeichnung', 'Psi-Wert'] or not all(headers):
            raise ValueError('Expected the original workbook headers in rows 2 and 3, columns A:J')
        records = []
        for cells in sheet.iter_rows(min_row=4, max_col=10):
            if not cells[0].value:
                continue
            record = {}
            for header, cell in zip(headers, cells):
                value = cell.value
                if isinstance(value, (datetime, date)):
                    value = value.strftime('%d.%m.%Y')
                record[header] = value
            cell = cells[0]
            size = cell.font.sz or 10
            record.update(
                row_color=_hex_color(cell.fill.fgColor, theme, palette, '#ffffff')
                if cell.fill.patternType == 'solid' else '#ffffff',
                text_size=int(size) if size == int(size) else size,
                text_color=_hex_color(cell.font.color, theme, palette, '#000000'),
            )
            records.append(record)
        return records
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--sheet', default='44er Wand')
    args = parser.parse_args()
    records = read_excel_rows(args.workbook, args.sheet)
    if not records:
        raise ValueError('No rows found in the selected worksheet')
    with args.output.open('w', encoding='utf-8', newline='') as output:
        writer = csv.DictWriter(output, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    print(f'Exported {len(records)} rows to {args.output}')


if __name__ == '__main__':
    main()
