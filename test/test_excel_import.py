from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Color
from berroutils.plugins.file_handler import JsonFileHandler

from src.thermal_bridge.excel_import import read_excel_rows
from src.thermal_bridge.psi_data import Psi


def test_excel_import_preserves_styles_and_hierarchy(tmp_path, monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '44er Wand'
    sheet.append(['Wärmebrücken'])
    sheet.append(['Bezeichnung', 'Psi-Wert', 'mit Referenzbauteil',
                  'berechnet durch', None, None, 'Datum', 'Nr.', 'BV', 'Name'])
    sheet.append([None, None, None, 'ebz', 'WP', 'VHAG'])
    sheet.append(['Außenwand auf Bodenplatte (EG)'])
    sheet['A4'].font = Font(size=12, color=Color(indexed=8))
    sheet.append(['Erhöhte Terrasse'])
    sheet['A5'].font = Font(size=10, color='FFFF0000')
    sheet.append(['AW44-P-15PPW2-160mm032', 0.0875, None, None, None,
                  'X', datetime(2017, 4, 13), 1142])
    sheet['A6'].font = Font(size=10, color='FF000000')
    sheet['A6'].fill = PatternFill('solid', fgColor='FFEAD1DC')
    source = tmp_path / 'source.xlsx'
    workbook.save(source)

    rows = read_excel_rows(source)
    assert rows[0]['text_size'] == 12
    assert rows[0]['text_color'] == '#000000'
    assert rows[0]['row_color'] == '#ffffff'
    assert rows[1]['text_color'] == '#ff0000'
    assert rows[2]['row_color'] == '#ead1dc'
    assert rows[2]['Datum'] == '13.04.2017'
    psi = Psi(filehandler=JsonFileHandler(file_path=tmp_path / 'data.json'))
    psi.update_from_excel(source)
    assert len(psi.data) == 1
    assert psi.data[0]['Waermebruecke'] == 'Außenwand auf Bodenplatte (EG)'
    assert psi.data[0]['Zusatzinfo Waermebruecke'] == 'Erhöhte Terrasse'
    assert psi.data[0]['Datum'] == '2017-04-13'

    sheet['B6'] = '0,.0209'
    workbook.save(source)
    answers = iter(['still invalid', '0.0209'])
    prompts = []

    def correct_value(prompt):
        prompts.append(prompt)
        return next(answers)

    monkeypatch.setattr('builtins.input', correct_value)
    psi.update_from_excel(source)
    assert len(psi.data) == 1
    assert psi.data[0]['Psi-Wert'] == 0.0209
    assert len(prompts) == 2
    assert '0,.0209' in prompts[0]
    assert 'AW44-P-15PPW2-160mm032' in prompts[0]
    assert 'still invalid' in prompts[1]
    assert read_excel_rows(source)[2]['Psi-Wert'] == '0,.0209'
